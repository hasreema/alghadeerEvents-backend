from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import List, Optional

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import openpyxl
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.models.event import Event


def _register_hebrew_font():
    if settings.hebrew_font_path:
        try:
            pdfmetrics.registerFont(TTFont('Hebrew', settings.hebrew_font_path))
            return 'Hebrew'
        except Exception:
            pass
    return 'Helvetica'


def generate_monthly_pdf(events: List[Event], month_label: str) -> bytes:
    buffer = BytesIO()
    font_name = _register_hebrew_font()

    p = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    p.setFont(font_name, 16)
    p.drawString(40, height - 50, f"דוח חודשי - {month_label}")

    p.setFont(font_name, 10)
    y = height - 90
    totals = {
        'revenue': Decimal('0'),
        'expenses': Decimal('0'),
        'labor': Decimal('0'),
        'profit': Decimal('0'),
    }

    for e in events:
        revenue = Decimal(str(e.payments_total or 0))
        expenses = Decimal(str(e.expenses_total or 0))
        labor = Decimal(str(e.labor_total or 0))
        profit = revenue - expenses - labor
        totals['revenue'] += revenue
        totals['expenses'] += expenses
        totals['labor'] += labor
        totals['profit'] += profit

        line = f"{e.date} | {e.title} | הכנסות {revenue} | הוצאות {expenses} | עבודה {labor} | רווח {profit}"
        p.drawString(40, y, line)
        y -= 18
        if y < 60:
            p.showPage()
            p.setFont(font_name, 10)
            y = height - 60

    p.setFont(font_name, 12)
    p.drawString(40, y - 10, f"סה""כ הכנסות: {totals['revenue']}")
    p.drawString(40, y - 28, f"סה""כ הוצאות: {totals['expenses']}")
    p.drawString(40, y - 46, f"סה""כ עבודה: {totals['labor']}")
    p.drawString(40, y - 64, f"סה""כ רווח: {totals['profit']}")

    p.showPage()
    p.save()
    buffer.seek(0)
    return buffer.read()


def export_events_excel(events: List[Event]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Events"

    headers = [
        "ID", "Title", "Date", "Status", "Payment Status", "Quoted", "Payments", "Expenses", "Labor", "Outstanding",
    ]
    ws.append(headers)

    for e in events:
        ws.append([
            e.id, e.title, str(e.date), e.status, e.payment_status,
            float(e.quoted_total or 0), float(e.payments_total or 0), float(e.expenses_total or 0), float(e.labor_total or 0), float(e.outstanding_amount or 0),
        ])

    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.read()